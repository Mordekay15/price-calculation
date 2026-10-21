"""The customer quote, pure (no Streamlit): its header, its lines (a price per
batch size for every part) and the files made from it: the quote PDF, the
Excel export and the clean DXFs for production."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, timedelta

from core.batches import BatchPrice
from core.work import BREAKDOWN


@dataclass(frozen=True)
class QuoteInfo:
    """The quote's header, as typed in the "Tarjous" section."""

    number: str = ""
    customer: str = ""
    customer_address: str = ""
    contact: str = ""
    seller: str = ""
    company: str = "Stremet Oy"
    company_address: str = ""
    quote_date: date = field(default_factory=date.today)
    valid_days: int = 30
    payment_terms: str = "14 pv netto"
    delivery: str = "Sisältyy"
    notes: str = ""

    @property
    def valid_until(self) -> date:
        return self.quote_date + timedelta(days=self.valid_days)


@dataclass(frozen=True)
class QuoteLine:
    """One part on the quote and its price at each batch size; or one
    assembly, with ``parts`` saying what goes into it (e.g. "Pohja × 1") and
    its prices ``core.assemblies.AssemblyPrice``."""

    name: str
    material: str = ""
    thickness: str = ""
    prices: tuple[BatchPrice, ...] = ()
    parts: tuple[str, ...] = ()

    @property
    def is_assembly(self) -> bool:
        return bool(self.parts)


def quote_totals(lines: list[QuoteLine], n_sizes: int) -> list[float | None]:
    """The quote's total (€, VAT 0 %) at each of ``n_sizes`` batch sizes; None
    where any line is unpriced."""
    totals = [[line.prices[i].total for line in lines] for i in range(n_sizes)]
    return [None if any(t is None for t in row) else sum(row) for row in totals]


def fmt_eur(value: float | None) -> str:
    """``1 234,56 €`` in the Finnish way; "" when unpriced."""
    if value is None:
        return ""
    return f"{value:,.2f} €".replace(",", " ").replace(".", ",")


def _fmt_date(d: date) -> str:
    return f"{d.day}.{d.month}.{d.year}"


# ── PDF ───────────────────────────────────────────────────────────────────────

def quote_pdf(info: QuoteInfo, lines: list[QuoteLine], sizes: list[int],
              totals: list[float | None], *, show_material: bool = True) -> bytes:
    """The quote as a PDF: the header, one block of rows per part (a row per
    batch size: quantity, unit price, sum) and the totals per batch size."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm,
                            title=f"Tarjous {info.number}".strip())
    styles = getSampleStyleSheet()
    small = styles["Normal"].clone("small", fontSize=9, leading=11)
    story = []

    head = Table([
        [Paragraph(f"<b>{_esc(info.company)}</b><br/>{_esc(info.company_address)}",
                   styles["Normal"]),
         Paragraph(f"<font size=18><b>Tarjous</b></font><br/>{_esc(info.number)}",
                   styles["Normal"])],
    ], colWidths=[100 * mm, 74 * mm])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                              ("ALIGN", (1, 0), (1, 0), "RIGHT")]))
    story += [head, Spacer(1, 6 * mm)]

    facts = [
        ("Tarjottu", "\n".join(filter(None, (info.customer, info.customer_address)))),
        ("Yhteyshenkilö", info.contact),
        ("Tarjouspäivä", _fmt_date(info.quote_date)),
        ("Tarjous voimassa", _fmt_date(info.valid_until)),
        ("Maksuehto", info.payment_terms),
        ("Toimitus", info.delivery),
        ("Myyjä", info.seller),
    ]
    facts_table = Table([[Paragraph(f"<b>{k}</b>", small), Paragraph(_esc(v, br=True), small)]
                         for k, v in facts if v], colWidths=[36 * mm, 138 * mm])
    facts_table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                     ("BOTTOMPADDING", (0, 0), (-1, -1), 1)]))
    story += [facts_table, Spacer(1, 6 * mm),
              Paragraph("<b>Tarjotut nimikkeet</b>", styles["Heading3"])]

    rows = [["Nimike", "Eräkoko", "Yksikköhinta", "Summa (ALV 0 %)"]]
    spans = []
    for line in lines:
        first = len(rows)
        name = line.name
        if show_material and (line.material or line.thickness):
            name += f"<br/><font size=8 color='#555555'>{_esc(line.material)} " \
                    f"{_esc(line.thickness)}{' mm' if line.thickness else ''}</font>"
        if line.parts:
            name += "<br/><font size=8 color='#555555'>Kokoonpano: " \
                    f"{_esc(', '.join(line.parts))}</font>"
        for i, p in enumerate(line.prices):
            rows.append([Paragraph(name, small) if i == 0 else "", f"{p.qty}",
                         fmt_eur(p.unit) or "–", fmt_eur(p.total) or "–"])
        if len(line.prices) > 1:
            spans.append(("SPAN", (0, first), (0, len(rows) - 1)))
        spans.append(("LINEBELOW", (0, len(rows) - 1), (-1, len(rows) - 1), 0.25,
                      colors.lightgrey))
    table = Table(rows, colWidths=[78 * mm, 24 * mm, 34 * mm, 38 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("LINEBELOW", (0, 0), (-1, 0), 0.75, colors.black),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        *spans,
    ]))
    story += [table, Spacer(1, 6 * mm)]

    total_rows = [["Yhteensä eräkoolla", "Summa (ALV 0 %)"]]
    total_rows += [[f"{q} kpl / nimike", fmt_eur(t) or "osa hinnoista puuttuu"]
                   for q, t in zip(sizes, totals)]
    tt = Table(total_rows, colWidths=[60 * mm, 50 * mm], hAlign="RIGHT")
    tt.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9),
                            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                            ("LINEBELOW", (0, 0), (-1, 0), 0.75, colors.black)]))
    story.append(tt)
    if info.notes:
        story += [Spacer(1, 6 * mm), Paragraph("<b>Lisätiedot</b>", small),
                  Paragraph(_esc(info.notes, br=True), small)]
    story += [Spacer(1, 4 * mm),
              Paragraph("Hinnat ilman arvonlisäveroa (ALV 0 %).", small)]
    doc.build(story)
    return buf.getvalue()


def _esc(text: str | None, br: bool = False) -> str:
    out = (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return out.replace("\n", "<br/>") if br else out


# ── Excel ─────────────────────────────────────────────────────────────────────

def quote_xlsx(info: QuoteInfo, lines: list[QuoteLine], sizes: list[int]) -> bytes:
    """The quote as an Excel workbook: "Tarjous" (a row per part and batch
    size) and "Erittely" (the cost breakdown per piece)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Tarjous"
    for k, v in (("Tarjous", info.number), ("Asiakas", info.customer),
                 ("Tarjouspäivä", _fmt_date(info.quote_date)),
                 ("Voimassa", _fmt_date(info.valid_until)),
                 ("Maksuehto", info.payment_terms), ("Toimitus", info.delivery),
                 ("Myyjä", info.seller)):
        ws.append([k, v])
    ws.append([])
    header = ["Nimike", "Materiaali", "Paksuus (mm)", "Eräkoko (kpl)",
              "Yksikköhinta (€)", "Summa (€, ALV 0 %)"]
    ws.append(header)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for line in lines:
        for p in line.prices:
            ws.append([line.name, line.material, line.thickness, p.qty,
                       _round(p.unit), _round(p.total)])

    ws2 = wb.create_sheet("Erittely")
    ws2.append(["Nimike", "Eräkoko (kpl)", "Materiaali €/kpl",
                *(f"{phase} €/kpl" for phase in BREAKDOWN), "Asetus ym. €/kpl",
                "Levyjä (ryhmä)", "Kpl-hinta €"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for line in lines:
        for p in line.prices:
            blank = ("",) * (len(BREAKDOWN) + 1)
            if line.is_assembly:
                ws2.append([line.name, p.qty, "", *blank, "", _round(p.unit)])
                continue
            w = p.work
            ws2.append([
                line.name, p.qty, _round(p.material),
                *((*map(_round, w.breakdown().values()), _round(w.batch_fixed / w.qty))
                  if w is not None else blank),
                p.sheets, _round(p.unit),
            ])
    sheets = [ws, ws2]
    assemblies = [line for line in lines if line.is_assembly]
    if assemblies:
        ws3 = wb.create_sheet("Kokoonpanot")
        sheets.append(ws3)
        ws3.append(["Kokoonpano", "Eräkoko (kpl)", "Osat €/kok.", "Ostetut osat €/kok.",
                    "Kokoonpanotyö €/kok.", "Yhteensä €/kok.", "Sisältö"])
        for cell in ws3[1]:
            cell.font = Font(bold=True)
        for line in assemblies:
            for p in line.prices:
                ws3.append([line.name, p.qty, _round(p.parts), _round(p.bought),
                            _round(p.work), _round(p.unit), ", ".join(line.parts)])
    for sheet in sheets:
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _round(value: float | None) -> float | str:
    return "" if value is None else round(value, 2)


# ── Clean DXFs for production ─────────────────────────────────────────────────

DEFAULT_DXF_NAME = "{part}_{material}_{thickness}mm"


@dataclass(frozen=True)
class CleanPart:
    """One part's cut geometry in mm: the outline, its holes, open cuts and
    (optional) markings such as bend lines."""

    part: str
    material: str
    thickness: str
    outline: list
    holes: list = field(default_factory=list)
    open_lines: list = field(default_factory=list)
    markings: list = field(default_factory=list)


def dxf_file_name(template: str, part: CleanPart, index: int) -> str:
    """The file name from the template: ``{part}``, ``{material}``,
    ``{thickness}`` and ``{index}``; unsafe characters become ``_``."""
    raw = (template or DEFAULT_DXF_NAME)
    for k, v in (("part", re.sub(r"\.(dxf|dwg)$", "", part.part, flags=re.I)),
                 ("material", part.material), ("thickness", part.thickness),
                 ("index", str(index))):
        raw = raw.replace("{" + k + "}", v or "")
    name = re.sub(r"[^\w.\-+]+", "_", raw).strip("_.") or f"osa_{index}"
    return name + ".dxf"


def clean_dxf(part: CleanPart, *, markings: bool = False) -> bytes:
    """One part as a clean DXF (R2010, mm): the outline and holes as closed
    polylines on layer ``CUT``, open cuts on ``CUT`` too, and the markings on
    ``MARKING`` when asked for. Nothing else: no frame, texts or dimensions."""
    import ezdxf

    doc = ezdxf.new("R2010", units=4)   # 4 = mm
    doc.layers.add("CUT", color=7)
    doc.layers.add("MARKING", color=3)
    msp = doc.modelspace()
    for loop in (part.outline, *part.holes):
        msp.add_lwpolyline([tuple(p) for p in loop], close=True, dxfattribs={"layer": "CUT"})
    for line in part.open_lines:
        msp.add_lwpolyline([tuple(p) for p in line], dxfattribs={"layer": "CUT"})
    if markings:
        for line in part.markings:
            msp.add_lwpolyline([tuple(p) for p in line], dxfattribs={"layer": "MARKING"})
    buf = io.StringIO()
    doc.write(buf)
    return buf.getvalue().encode("utf-8")


def clean_dxf_zip(parts: list[CleanPart], template: str = DEFAULT_DXF_NAME, *,
                  markings: bool = False) -> bytes:
    """Every part's clean DXF in one ZIP; repeated names get ``_2``, ``_3``…"""
    buf = io.BytesIO()
    used: dict[str, int] = {}
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for i, part in enumerate(parts, start=1):
            name = dxf_file_name(template, part, i)
            n = used.get(name, 0) + 1
            used[name] = n
            if n > 1:
                name = name[:-4] + f"_{n}.dxf"
            zf.writestr(name, clean_dxf(part, markings=markings))
    return buf.getvalue()
