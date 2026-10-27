"""The quote's files: PDF, Excel and the clean DXFs."""

import io
import zipfile
from datetime import date

import ezdxf
import openpyxl

from core.batches import BatchPrice
from core.work import WorkCost
from core.quote import (
    CleanPart,
    QuoteInfo,
    QuoteLine,
    clean_dxf,
    clean_dxf_zip,
    dxf_file_name,
    fmt_eur,
    quote_pdf,
    quote_xlsx,
)

SQUARE = [(0, 0), (100, 0), (100, 50), (0, 50)]
HOLE = [(10, 10), (20, 10), (20, 20), (10, 20)]


def _info():
    return QuoteInfo(number="T-1", customer="Asiakas Oy", quote_date=date(2026, 10, 1),
                     notes="Rivi 1\nRivi 2 <b>")


def _lines():
    prices = (BatchPrice(1, 10.0, WorkCost(1, cutting=20.0), 1),
              BatchPrice(10, 5.0, WorkCost(10, cutting=2.0, setup=10.0), 1))
    return [QuoteLine("osa_1", "S355MC", "2", prices),
            QuoteLine("osa_2", "RST 2B", "1.5", (BatchPrice(1, None, WorkCost(1), 0),
                                                  BatchPrice(10, None, WorkCost(10), 0)))]


def test_fmt_eur_is_finnish():
    assert fmt_eur(1234.5) == "1 234,50 €"
    assert fmt_eur(None) == ""


def test_valid_until():
    assert _info().valid_until == date(2026, 10, 31)


def test_line_prices():
    a, b = _lines()[0].prices
    assert a.unit == 30.0
    assert b.unit == 8.0 and b.total == 80.0
    assert _lines()[1].prices[0].unit is None


def test_quote_pdf_is_a_pdf():
    data = quote_pdf(_info(), _lines(), [1, 10], [30.0, None])
    assert data.startswith(b"%PDF")
    assert quote_pdf(_info(), _lines(), [1, 10], [30.0, None],
                     show_material=False).startswith(b"%PDF")


def test_the_quote_can_be_made_with_the_order_quantities():
    assert quote_pdf(_info(), _lines(), ["Tilausmäärillä"], [30.0],
                     qty_label="Määrä").startswith(b"%PDF")
    wb = openpyxl.load_workbook(io.BytesIO(quote_xlsx(_info(), _lines(), qty_label="Määrä")))
    assert "Määrä (kpl)" in [c.value for c in wb["Erittely"][1]]


def test_quote_xlsx_opens_with_both_sheets():
    wb = openpyxl.load_workbook(io.BytesIO(quote_xlsx(_info(), _lines(), [1, 10])))
    assert wb.sheetnames == ["Tarjous", "Erittely"]
    values = [c.value for row in wb["Tarjous"].iter_rows() for c in row]
    assert "osa_1" in values and "Asiakas Oy" in values


def test_dxf_file_name_template():
    p = CleanPart("bracket/L.dxf", "S355MC", "2", SQUARE)
    assert dxf_file_name("{part}_{material}_{thickness}mm", p, 1) == "bracket_L_S355MC_2mm.dxf"
    assert dxf_file_name("{index}-{part}", p, 7) == "7-bracket_L.dxf"
    assert dxf_file_name("", p, 1) == "bracket_L_S355MC_2mm.dxf"


def test_clean_dxf_layers():
    p = CleanPart("a", "S355MC", "2", SQUARE, holes=[HOLE],
                  markings=[[(0, 25), (100, 25)]])
    doc = ezdxf.read(io.StringIO(clean_dxf(p).decode()))
    layers = [e.dxf.layer for e in doc.modelspace()]
    assert layers == ["CUT", "CUT"]
    doc = ezdxf.read(io.StringIO(clean_dxf(p, markings=True).decode()))
    assert [e.dxf.layer for e in doc.modelspace()] == ["CUT", "CUT", "MARKING"]
    assert doc.header["$INSUNITS"] == 4


def test_clean_dxf_zip_deduplicates_names():
    parts = [CleanPart("a.dxf", "S355MC", "2", SQUARE)] * 2 + [
        CleanPart("b.dxf", "S355MC", "2", SQUARE)]
    with zipfile.ZipFile(io.BytesIO(clean_dxf_zip(parts))) as zf:
        assert zf.namelist() == ["a_S355MC_2mm.dxf", "a_S355MC_2mm_2.dxf", "b_S355MC_2mm.dxf"]
